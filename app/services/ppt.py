from __future__ import annotations

from pathlib import Path
from typing import Dict

from pptx import Presentation

from app.models import Profile

SHAPE_NAMES = [
    "TITLE_COMPANY_NAME",
    "TITLE_LOGO",
    "BO_OVERALL",
    "BO_SEGMENTS",
    "BO_CLIENTS",
    "BO_INDUSTRIAL",
    "FIN_TABLE",
    "MAP_IMAGE",
    "OWNERSHIP_TABLE",
    "DEALS_TABLE",
    "CITATIONS_FOOTER",
]


def _safe_text(value: str | None) -> str:
    return value if value else "Not publicly disclosed"


def compose_ppt(profile: Profile, export_dir: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    textbox_map: Dict[str, str] = {
        "TITLE_COMPANY_NAME": profile.company_input.name,
        "BO_OVERALL": _safe_text(profile.business_overview.overall.value),
        "BO_SEGMENTS": _safe_text(profile.business_overview.segments.value),
        "BO_CLIENTS": _safe_text(profile.business_overview.clients.value),
        "BO_INDUSTRIAL": _safe_text(profile.business_overview.industrial_setup.value),
        "FIN_TABLE": (
            f"Revenue: {_safe_text(str(profile.financials.revenue.value) if profile.financials.revenue.value is not None else None)}\n"
            f"EBITDA: {_safe_text(str(profile.financials.ebitda.value) if profile.financials.ebitda.value is not None else None)}\n"
            f"Market cap: {_safe_text(str(profile.financials.market_cap.value) if profile.financials.market_cap.value is not None else None)}\n"
            f"Employees: {_safe_text(str(profile.financials.employees.value) if profile.financials.employees.value is not None else None)}"
        ),
        "OWNERSHIP_TABLE": _safe_text(
            "\n".join(
                [f"{e.name} - {e.stake if e.stake is not None else 'Not publicly disclosed'}" for e in profile.ownership.entries]
            )
            if profile.ownership.entries
            else profile.ownership.note
        ),
        "DEALS_TABLE": _safe_text(
            "\n".join([f"{d.year} {d.type}: {d.counterparty}" for d in profile.transactions])
            if profile.transactions
            else None
        ),
        "CITATIONS_FOOTER": " ".join(sorted(profile.evidence.keys())) if profile.evidence else "Not publicly disclosed",
    }

    left = 100000
    top = 100000
    width = 4000000
    height = 300000

    for shape_name in SHAPE_NAMES:
        textbox = slide.shapes.add_textbox(left, top, width, height)
        textbox.name = shape_name
        textbox.text = textbox_map.get(shape_name, "")
        top += 320000

    export_dir.mkdir(parents=True, exist_ok=True)
    output = export_dir / f"{profile.profile_id}.pptx"
    prs.save(output)
    return output
